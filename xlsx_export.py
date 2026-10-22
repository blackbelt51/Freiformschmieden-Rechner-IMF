# -*- coding: utf-8 -*-
"""
Excel-Export des Stichplans (openpyxl)
======================================
Blätter:
  Stichplan      druckfertige Übersicht: Kopfdaten, Kennzahlen, Stichtabelle, Hinweise, Diagramme
  Details        alle berechneten Größen je Stich
  Parameter      sämtliche Eingaben, Konstanten und das Rechenmodell
  Diagrammdaten  Zeitreihe und Hilfswerte, auf die sich die Diagramme beziehen

Alle Zahlen sind Ergebnisse des Programms und als Werte hinterlegt (keine Formeln):
Temperaturen und Kräfte stammen aus der Simulation und lassen sich in Excel nicht nachrechnen.
"""
import math

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference, ScatterChart, Series
from openpyxl.chart.marker import DataPoint
from openpyxl.chart.series import SeriesLabel
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText, Text
from openpyxl.chart.title import Title
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import CharacterProperties, Font as DFont, Paragraph, ParagraphProperties, RegularTextRun
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.pagebreak import Break

FONT = "Arial"
INK, MUTED = "1C2833", "5A6774"
LINE, LINE_STRONG = "C6CDD3", "8F9BA6"
PANEL, PANEL2, ZEBRA = "EDF0F2", "E1E6EA", "F6F8F9"
COLD, COLD_SOFT = "2D6AA3", "D6E2EE"
HEAT, HEAT2, HEAT3 = "D24E17", "EE9B3E", "A92A12"
WARN, WARN_BG = "8A5200", "FBEFD6"
ERR, ERR_BG = "B3261E", "FBE6E4"
INFO_BG = "DCE7F1"
INPUT = "1F5FBF"          # Eingaben (wie in Excel-Modellen üblich blau)

SHAPES = {"rund": ("Rund", "Ø"), "vier": ("Vierkant", "□"), "acht": ("Achtkant", "SW")}


# ---------------------------------------------------------------- kleine Helfer
def de(x, d=0):
    """Zahl als Text mit Dezimalkomma und Tausenderpunkt (für Fließtext in Zellen)."""
    if x is None or (isinstance(x, float) and not math.isfinite(x)):
        return "–"
    s = f"{abs(x):,.{d}f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return ("−" if x < 0 and round(abs(x), d) != 0 else "") + s


def dim_txt(x):
    return de(x, 0 if abs(x - round(x)) < 0.05 else 1)


def geo_txt(shape, d, L):
    name, pre = SHAPES[shape]
    return f"{name} {pre} {dim_txt(d)} × {de(L, 0)} mm"


def font(size=10, bold=False, color=INK, italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def fill(color):
    return PatternFill("solid", start_color=color, end_color=color)


def side(style="thin", color=LINE):
    return Side(style=style, color=color)


def put(ws, ref, value, f=None, fmt=None, al=None, fl=None, br=None):
    c = ws[ref] if isinstance(ref, str) else ws.cell(row=ref[0], column=ref[1])
    c.value = value.strip() if isinstance(value, str) else value
    c.font = f or font()
    if fmt:
        c.number_format = fmt
    if al:
        c.alignment = al
    if fl:
        c.fill = fl
    if br:
        c.border = br
    return c


def merge_put(ws, r1, c1, r2, c2, value, **kw):
    if (r1, c1) != (r2, c2):
        ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
    cell = put(ws, (r1, c1), value, **kw)
    # Rahmen und Füllung auf alle Zellen des Bereichs übertragen (sonst fehlen sie bei verbundenen Zellen)
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            if (r, c) == (r1, c1):
                continue
            if kw.get("fl"):
                ws.cell(row=r, column=c).fill = kw["fl"]
            if kw.get("br"):
                ws.cell(row=r, column=c).border = kw["br"]
    return cell


def setup_print(ws, landscape=True, last_col=None, last_row=None, title=""):
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.6, bottom=0.7, header=0.3, footer=0.3)
    ws.print_options.horizontalCentered = True
    if last_col and last_row:
        ws.print_area = f"A1:{get_column_letter(last_col)}{last_row}"
    ws.oddFooter.left.text = title
    ws.oddFooter.left.size = 8
    ws.oddFooter.left.font = f"{FONT},Regular"
    ws.oddFooter.right.text = "Seite &P von &N"
    ws.oddFooter.right.size = 8
    ws.oddFooter.right.font = f"{FONT},Regular"


def _bite_txt(proc):
    mode = proc.get("bite_mode", "min")
    if mode == "const" and proc.get("sB_const"):
        return f"s_B = {de(min(proc['sB_const'], proc['b_die']), 0)} mm"
    if mode == "sbh" and proc.get("sB_rel"):
        return f"s_B/h = {de(proc['sB_rel'], 2)}"
    return "s_B = Sattelbreite"


# ---------------------------------------------------------------- Diagramm-Formatierung
def _cp(size=900, bold=False, color=MUTED):
    return CharacterProperties(latin=DFont(typeface=FONT), sz=size, b=bold, solidFill=color)


def _rich(size=900, bold=False, color=MUTED):
    cp = _cp(size, bold, color)
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def _title(text, size=1100, bold=True, color=INK):
    cp = _cp(size, bold, color)
    para = Paragraph(pPr=ParagraphProperties(defRPr=cp), r=[RegularTextRun(rPr=cp, t=text)])
    return Title(tx=Text(rich=RichText(p=[para])), overlay=False)


def _style_axes(ch, x_title, y_title):
    for ax, t in ((ch.x_axis, x_title), (ch.y_axis, y_title)):
        ax.title = _title(t, 900, False, MUTED)
        ax.txPr = _rich(900, False, MUTED)
        ax.delete = False
        ax.spPr = GraphicalProperties(ln=LineProperties(solidFill=LINE_STRONG))
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill=LINE))
    ch.legend.position = "b"
    ch.legend.txPr = _rich(900, False, INK)
    try:
        ch.graphical_properties = GraphicalProperties(ln=LineProperties(noFill=True))
    except Exception:
        pass


def _nice(lo, hi, n=6):
    span = max(hi - lo, 1e-9)
    raw = span / n
    p = 10 ** math.floor(math.log10(raw))
    f = raw / p
    step = (1 if f < 1.5 else 2 if f < 3 else 5 if f < 7 else 10) * p
    return math.floor(lo / step) * step, math.ceil(hi / step) * step, step


def _decimate(t, y, buckets=350):
    """Ausdünnen unter Erhalt der Extremwerte (die Oberflächentemperatur hat kurze Einbrüche beim Kontakt)."""
    n = len(t)
    if n <= 2 * buckets:
        return list(range(n))
    keep = {0, n - 1}
    size = n / buckets
    for k in range(buckets):
        lo, hi = int(k * size), min(n, int((k + 1) * size))
        if hi <= lo:
            continue
        seg = range(lo, hi)
        keep.add(min(seg, key=lambda i: y[i]))
        keep.add(max(seg, key=lambda i: y[i]))
    return sorted(keep)


# ---------------------------------------------------------------- Export
def export_stichplan(path, D):
    """
    D: dict mit
      version, created (datetime), mat(name, rho, cp, k), block(shape, d, r, L, A, mass),
      proc (Prozessparameter des Rechenkerns inkl. F_max, T_min, T_max), kf (Modell), start (T_Kern, T_Ofl, T_Mittel),
      rows (je Stich: Geometrie + Ergebnis + Eingabeflags), series(t, core, surf, contacts), notes [(art, text)],
      model [(Überschrift, Text)]
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Stichplan"
    wd = wb.create_sheet("Details")
    wp = wb.create_sheet("Parameter")
    wc = wb.create_sheet("Diagrammdaten")
    mat, blk, proc, rows = D["mat"], D["block"], D["proc"], D["rows"]
    title = f"Stichplan {mat['name']}"
    stamp = D["created"].strftime("%d.%m.%Y, %H:%M Uhr")
    wb.properties.title = title
    wb.properties.creator = f"Freiformschmieden-Rechner {D['version']}"
    wb.properties.subject = "Stichplan Freiformschmieden"

    chart_refs = _sheet_chartdata(wc, D)
    table_rows = _sheet_overview(ws, D, title, stamp, chart_refs)
    _sheet_details(wd, D, title)
    _sheet_parameter(wp, D, title)
    ws.sheet_properties.tabColor = COLD
    wd.sheet_properties.tabColor = LINE_STRONG
    wp.sheet_properties.tabColor = LINE_STRONG
    wc.sheet_properties.tabColor = PANEL2
    wb.active = 0
    wb.save(path)
    return table_rows


# ---------------------------------------------------------------- Blatt 1: Übersicht
OV_COLS = [
    # key, Überschrift, Breite, Zahlenformat
    ("nr", "Stich", 6.5, "0"),
    ("form", "Form", 11, None),
    ("d", "Maß D/a/SW\n[mm]", 10, "0.0"),
    ("L", "Länge\n[mm]", 9, "#,##0"),
    ("ratio", "Verschmie-\ndungsgrad", 10.5, "0.00"),
    ("phi", "φ\n[–]", 7.5, "0.000"),
    ("n_rot", "Dreh-\nlagen", 7, "0"),
    ("n_b", "Bisse", 8.5, "0"),
    ("strokes", "Hübe", 7, "0"),
    ("sB", "Bisslänge\ns_B [mm]", 9.5, "0"),
    ("bc", "Kontaktbreite\nb_K [mm]", 12, "0"),
    ("t_c", "Kontaktzeit\nt_K [s]", 10.5, "0.00"),
    ("t_end_min", "Ende nach\n[min]", 9.5, "0.0"),
    ("kf", "k_f\n[MPa]", 7.5, "0.0"),
    ("F", "Presskraft\n[MN]", 10, "0.00"),
    ("T_core", "Kern\n[°C]", 7.5, "0"),
    ("T_surf", "Oberfläche\n[°C]", 10.5, "0"),
    ("T_contact", "Kontaktzone\nmin. [°C]", 11.5, "0"),
]
OV_GROUPS = [("Geometrie", 2, 6), ("Ablauf an der Presse", 7, 13), ("Kraft", 14, 15), ("Temperatur", 16, 18)]


def _flags(r, proc):
    """Zellmarkierungen für Grenzwertüberschreitungen: key -> 'err' | 'warn'."""
    fl = {}
    if proc is None:
        return fl
    if r.get("F") is not None and r["F"] > proc["F_max"]:
        fl["F"] = "err"
    if r.get("T_core") is not None and r["T_core"] > proc["T_max"]:
        fl["T_core"] = "err"
    if r.get("T_surf") is not None and r["T_surf"] < proc["T_min"]:
        fl["T_surf"] = "warn"
    if r.get("T_contact") is not None and r["T_contact"] < proc["T_min"]:
        fl["T_contact"] = "warn"
    return fl


def _sheet_overview(ws, D, title, stamp, refs):
    mat, blk, proc, rows = D["mat"], D["block"], D["proc"], D["rows"]
    ncol = len(OV_COLS)
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 100
    for j, (_k, _h, w, _f) in enumerate(OV_COLS, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w

    # --- Kopf
    put(ws, "A1", title, font(18, True))
    ws.row_dimensions[1].height = 30
    proj = f"Projekt „{D['project']}“, " if D.get("project") else ""
    put(ws, "A2", f"{proj}Freiformschmieden, Recken mit Flachsätteln. Erstellt am {stamp} mit dem "
                  f"Freiformschmieden-Rechner {D['version']}.", font(9.5, color=MUTED))
    for j in range(1, ncol + 1):
        ws.cell(row=2, column=j).border = Border(bottom=side("medium", INK))
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 8

    # --- Kopfdaten (drei Spalten mit Bezeichnung / Wert)
    last = rows[-1]
    kf = D["kf"]
    kf_txt = "Sellars-Tegart" if kf["type"] == "st" else "Hensel-Spittel"
    info = [
        [("Werkstoff", mat["name"]),
         ("Rohblock", f"{geo_txt(blk['shape'], blk['d'], blk['L'])}, {de(blk['mass'], 1)} kg"),
         ("Schmiedestück", geo_txt(last["shape"], last["d"], last["L"]))],
        [("Ofenentnahme", f"{de(proc['T_start'], 0)} °C, Transport {de(proc['t_transfer'], 0)} s"),
         ("Presse, Sättel", f"v = {de(proc['v'], 0)} mm/s, {_bite_txt(proc)}, "
                           f"Sättel {de(proc['b_die'], 0)} mm, {de(proc['T_die'], 0)} °C"),
         ("Fließspannung", kf_txt)],
        [("Wärmeübergang", f"h_K {de(proc['h_c'], 0)}, Luft {de(proc['h_air'], 0)} W/(m² K)"),
         ("Grenzen", f"F ≤ {de(proc['F_max'], 1)} MN, Oberfläche ≥ {de(proc['T_min'], 0)} °C, "
                     f"Kern ≤ {de(proc['T_max'], 0)} °C"),
         ("Umformung", f"η_F = {de(proc['eta_F'], 2)}, μ = {de(proc['mu'], 2)}")],
    ]
    spans = [((1, 2), (3, 6)), ((7, 8), (9, 13)), ((14, 15), (16, 18))]
    for i, line in enumerate(info):
        r = 4 + i
        ws.row_dimensions[r].height = 17
        for (lab, val), ((l1, l2), (v1, v2)) in zip(line, spans):
            merge_put(ws, r, l1, r, l2, lab, f=font(9, color=MUTED), al=Alignment(vertical="center"))
            merge_put(ws, r, v1, r, v2, val, f=font(10), al=Alignment(vertical="center"))
    ws.row_dimensions[7].height = 10

    # --- Kennzahlen
    total_phi = sum(r["phi"] for r in rows)
    F_max = max(r["F"] for r in rows)
    t_min = last["t_end"] / 60.0
    tiles = [
        ("Verschmiedungsgrad", math.exp(total_phi), '0.00" : 1"', COLD, INK),
        ("Umformgrad φ gesamt", total_phi, "0.00", COLD, INK),
        ("Hübe gesamt", sum(r["strokes"] for r in rows), "0", COLD, INK),
        ("Dauer ab Ofenentnahme", t_min, '0" min"' if t_min >= 10 else '0.0" min"', COLD, INK),
        ("Kern am Ende", last["T_core"], '0" °C"', HEAT, HEAT),
        ("Presskraft max.", F_max, '0.00" MN"', ERR if F_max > proc["F_max"] else COLD,
         ERR if F_max > proc["F_max"] else INK),
    ]
    r_lab, r_val = 8, 9
    ws.row_dimensions[r_lab].height = 18
    ws.row_dimensions[r_val].height = 32
    for i, (lab, val, fmt, accent, vcol) in enumerate(tiles):
        c1, c2 = 1 + 3 * i, 3 + 3 * i
        b_first = Border(left=side("thick", accent))
        merge_put(ws, r_lab, c1, r_lab, c2, lab, f=font(9, color=MUTED), fl=fill(PANEL),
                  al=Alignment(horizontal="left", vertical="bottom", indent=1))
        merge_put(ws, r_val, c1, r_val, c2, val, f=font(18, True, vcol), fmt=fmt, fl=fill(PANEL),
                  al=Alignment(horizontal="left", vertical="center", indent=1))
        ws.cell(row=r_lab, column=c1).border = b_first
        ws.cell(row=r_val, column=c1).border = b_first
        # schmale Lücke zwischen den Kacheln: rechter Rand in Hintergrundfarbe
        for r in (r_lab, r_val):
            cell = ws.cell(row=r, column=c2)
            cell.border = Border(left=cell.border.left, right=side("thick", "FFFFFF"))
    ws.row_dimensions[10].height = 12

    # --- Stichtabelle
    r_grp, r_head = 11, 12
    for name, c1, c2 in OV_GROUPS:
        merge_put(ws, r_grp, c1, r_grp, c2, name, f=font(9, True, MUTED),
                  al=Alignment(horizontal="center", vertical="bottom"),
                  br=Border(bottom=side("thin", LINE_STRONG)))
    ws.row_dimensions[r_grp].height = 16
    for j, (_k, head, _w, _f) in enumerate(OV_COLS, start=1):
        put(ws, (r_head, j), head, font(9, True), al=Alignment(horizontal="center", vertical="center", wrap_text=True),
            fl=fill(PANEL2), br=Border(bottom=side("medium", INK)))
    ws.row_dimensions[r_head].height = 30

    def body_row(r, vals, inputs=(), flags=None, zebra=False, start_row=False):
        flags = flags or {}
        for j, (key, _h, _w, fmt) in enumerate(OV_COLS, start=1):
            v = vals.get(key)
            col = INPUT if key in inputs else (MUTED if start_row else INK)
            bg = ZEBRA if zebra else None
            bold = False
            if key in flags:
                kind = flags[key]
                col, bg, bold = (ERR, ERR_BG, True) if kind == "err" else (WARN, WARN_BG, True)
            al = Alignment(horizontal="center" if key == "nr" else ("left" if key == "form" else "right"),
                           vertical="center", indent=1 if key == "form" else 0)
            put(ws, (r, j), v, font(10, bold or key == "nr", col), fmt=fmt if v is not None else None, al=al,
                fl=fill(bg) if bg else None, br=Border(bottom=side("thin", LINE)))
        ws.row_dimensions[r].height = 19

    r = r_head + 1
    st = D["start"]
    body_row(r, dict(nr=0, form=SHAPES[blk["shape"]][0], d=blk["d"], L=blk["L"], T_core=st[0], T_surf=st[1],
                     t_end_min=D.get("t_transfer", 0.0) / 60.0),
             inputs=("form", "d", "L"), start_row=True)
    first_pass_row = r + 1
    for i, row in enumerate(rows):
        r += 1
        vals = dict(row)
        vals["form"] = SHAPES[row["shape"]][0]
        vals["t_end_min"] = row["t_end"] / 60.0
        inputs = {"form", "d" if row["mode"] == "dim" else "L"}
        if row.get("nrot_manual"):
            inputs.add("n_rot")
        if row.get("bk_manual"):
            inputs.add("bc")
        body_row(r, vals, inputs=inputs, flags=_flags(row, proc), zebra=(i % 2 == 1))
    last_pass_row = r
    for j in range(1, ncol + 1):
        c = ws.cell(row=r, column=j)
        c.border = Border(bottom=side("medium", INK))

    # --- Erläuterungen
    r += 1
    ws.row_dimensions[r].height = 6
    r += 1
    legend = ("Blau: Eingabe im Programm, schwarz: berechnet. Stich 0 ist der Rohblock nach dem Transport zur Presse. "
              "Maß: Durchmesser D (Rund), Kantenlänge a (Vierkant) oder Schlüsselweite SW (Achtkant). Kern und "
              "Oberfläche am Stichende, Oberfläche als Mittel über den Umfang; Kontaktzone: niedrigste Temperatur der "
              "Kontaktfläche am Hubende. Rot: Grenzwert überschritten, orange: Mindesttemperatur unterschritten. "
              "Alle Werte sind Ergebnisse des Programms; Änderungen in dieser Datei werden nicht nachgerechnet.")
    merge_put(ws, r, 1, r, ncol, legend, f=font(8.5, color=MUTED), al=Alignment(wrap_text=True, vertical="top"))
    ws.row_dimensions[r].height = 36

    # --- Hinweise
    notes = [(k, t) for k, t in D["notes"] if k in ("err", "warn")]
    if not notes:
        notes = [("ok", "Alle Grenzwerte werden eingehalten.")]
    r += 1
    ws.row_dimensions[r].height = 8
    for kind, text in notes:
        r += 1
        bg, bar = {"err": (ERR_BG, ERR), "warn": (WARN_BG, WARN), "ok": (INFO_BG, COLD)}[kind]
        merge_put(ws, r, 1, r, ncol, text, f=font(10), fl=fill(bg),
                  al=Alignment(horizontal="left", wrap_text=True, vertical="center", indent=1))
        ws.cell(row=r, column=1).border = Border(left=side("thick", bar))
        ws.row_dimensions[r].height = 18 if len(text) < 160 else 32

    # --- Diagramme (eigene Druckseite)
    r += 2
    chart_row = r
    ws.row_breaks.append(Break(id=chart_row - 1))
    _chart_temperature(ws, refs, D, f"A{chart_row}")
    _chart_force(ws, refs, D, f"M{chart_row}", first_pass_row, last_pass_row)
    last_row = chart_row + 23
    setup_print(ws, True, ncol, last_row, title)
    ws.print_title_rows = f"{r_grp}:{r_head}" if len(rows) > 12 else None
    return first_pass_row, last_pass_row


def _chart_temperature(ws, refs, D, anchor):
    wc, r0, r1 = refs["sheet"], refs["t_first"], refs["t_last"]
    proc = D["proc"]
    ch = ScatterChart()
    ch.title = _title(f"Temperaturverlauf {D['mat']['name']}", 1100, True, INK)
    ch.style = 2
    ch.scatterStyle = "line"
    ch.height, ch.width = 11.5, 21.5
    xs = Reference(wc, min_col=2, min_row=r0, max_row=r1)
    for col, color, width in ((3, HEAT, 2.25), (4, HEAT2, 1.5)):
        s = Series(Reference(wc, min_col=col, min_row=r0 - 1, max_row=r1), xs, title_from_data=True)
        s.marker.symbol = "none"
        s.smooth = False
        s.graphicalProperties.line.solidFill = color
        s.graphicalProperties.line.width = int(width * 12700)
        ch.series.append(s)
    c0, c1 = refs["c_first"], refs["c_last"]
    if c1 >= c0:
        s = Series(Reference(wc, min_col=7, min_row=c0 - 1, max_row=c1),
                   Reference(wc, min_col=6, min_row=c0, max_row=c1), title_from_data=True)
        s.marker.symbol = "circle"
        s.marker.size = 5
        s.marker.graphicalProperties = GraphicalProperties(solidFill=HEAT3, ln=LineProperties(solidFill=HEAT3))
        s.graphicalProperties.line.noFill = True
        ch.series.append(s)
    for col in (10, 11):
        s = Series(Reference(wc, min_col=col, min_row=refs["g_first"] - 1, max_row=refs["g_first"] + 1),
                   Reference(wc, min_col=9, min_row=refs["g_first"], max_row=refs["g_first"] + 1),
                   title_from_data=True)
        s.marker.symbol = "none"
        s.graphicalProperties.line.solidFill = COLD
        s.graphicalProperties.line.width = int(1.25 * 12700)
        s.graphicalProperties.line.dashStyle = "dash"
        ch.series.append(s)
    _style_axes(ch, "Zeit ab Ofenentnahme in min", "Temperatur in °C")
    lo, hi, step = _nice(refs["T_lo"], refs["T_hi"], 6)
    ch.y_axis.scaling.min, ch.y_axis.scaling.max, ch.y_axis.majorUnit = lo, hi, step
    x_lo, x_hi, x_step = _nice(0.0, refs["t_max"], 8)
    ch.x_axis.scaling.min, ch.x_axis.scaling.max, ch.x_axis.majorUnit = 0, x_hi, x_step
    ch.x_axis.number_format = "0"
    ch.y_axis.number_format = "0"
    ch.x_axis.majorGridlines = None
    ws.add_chart(ch, anchor)


def _chart_force(ws, refs, D, anchor, r_first, r_last):
    proc = D["proc"]
    rows = D["rows"]
    F_peak = max(r["F"] for r in rows)
    show_limit = proc["F_max"] <= 2.5 * F_peak
    bar = BarChart()
    bar.type = "col"
    bar.style = 2
    bar.height, bar.width = 11.5, 11.0
    t = "Presskraft je Stich" + ("" if show_limit else f" (Grenze {de(proc['F_max'], 1)} MN)")
    bar.title = _title(t, 1100, True, INK)
    data = Reference(ws, min_col=15, min_row=r_first, max_row=r_last)
    cats = Reference(ws, min_col=1, min_row=r_first, max_row=r_last)
    bar.add_data(data, titles_from_data=False)
    bar.set_categories(cats)
    s = bar.series[0]
    s.tx = SeriesLabel(v="Presskraft")
    s.graphicalProperties = GraphicalProperties(solidFill=COLD, ln=LineProperties(noFill=True))
    for i, r in enumerate(rows):
        if r["F"] > proc["F_max"]:
            pt = DataPoint(idx=i)
            pt.graphicalProperties = GraphicalProperties(solidFill=ERR, ln=LineProperties(noFill=True))
            s.dPt.append(pt)
    bar.gapWidth = 60
    _style_axes(bar, "Stich", "Presskraft in MN")
    bar.y_axis.scaling.min = 0
    top = max(F_peak, proc["F_max"] if show_limit else 0.0) * 1.12
    _lo, hi, step = _nice(0.0, top, 5)
    bar.y_axis.scaling.max, bar.y_axis.majorUnit = hi, step
    bar.y_axis.number_format = "0.0" if hi < 10 else "0"
    if show_limit:
        wc = refs["sheet"]
        line = LineChart()
        line.add_data(Reference(wc, min_col=15, min_row=refs["f_first"] - 1, max_row=refs["f_last"]),
                      titles_from_data=True)
        ls = line.series[0]
        ls.marker.symbol = "none"
        ls.smooth = False
        ls.graphicalProperties.line.solidFill = ERR
        ls.graphicalProperties.line.width = int(1.5 * 12700)
        ls.graphicalProperties.line.dashStyle = "dash"
        bar += line
        bar.legend.position = "b"
    else:
        bar.legend = None
    ws.add_chart(bar, anchor)


# ---------------------------------------------------------------- Blatt 2: Details
DET_COLS = [
    # key, Überschrift, Breite, Format, Gruppe
    ("nr", "Stich", 6.5, "0"),
    ("form", "Form", 11, None),
    ("d", "Maß\n[mm]", 8.5, "0.0"),
    ("L", "Länge\n[mm]", 8.5, "#,##0"),
    ("A_cm2", "Querschnitt\n[cm²]", 11, "0.0"),
    ("ratio", "Verschmie-\ndungsgrad", 10.5, "0.000"),
    ("phi", "φ Stich\n[–]", 8, "0.000"),
    ("phi_cum", "φ gesamt\n[–]", 9, "0.000"),
    ("n_rot", "Dreh-\nlagen", 7, "0"),
    ("n_b", "Bisse", 7, "0"),
    ("strokes", "Hübe", 7, "0"),
    ("sB", "Bisslänge\ns_B [mm]", 9.5, "0.0"),
    ("feed", "Vorschub\n[mm]", 9.5, "0.0"),
    ("sB_h", "s_B/h\n[–]", 7, "0.00"),
    ("bc", "Kontakt-\nbreite [mm]", 10, "0.0"),
    ("A_K_cm2", "Kontaktfläche\nje Sattel [cm²]", 13, "0.0"),
    ("dh", "Hubweg\n[mm]", 8.5, "0.0"),
    ("t_c", "Kontaktzeit\n[s]", 10.5, "0.000"),
    ("phid", "φ̇\n[1/s]", 7.5, "0.000"),
    ("kf", "k_f\n[MPa]", 7.5, "0.0"),
    ("F", "Presskraft\n[MN]", 10, "0.000"),
    ("t_pass", "Stichdauer\n[s]", 10.5, "0"),
    ("t_end_min", "Stichende\n[min]", 9.5, "0.00"),
    ("dT_def", "Umform-\nwärme [K]", 9.5, "+0.0;−0.0;0.0"),
    ("dT_die_s", "Sattel-\nkontakt [K]", 10, "+0.0;−0.0;0.0"),
    ("dT_air", "Luft\n[K]", 7.5, "+0.0;−0.0;0.0"),
    ("dT_mean", "Summe\n[K]", 7.5, "+0.0;−0.0;0.0"),
    ("T_core", "Kern\n[°C]", 7.5, "0.0"),
    ("T_surf", "Oberfläche\n[°C]", 10.5, "0.0"),
    ("T_mean", "Mittel\n[°C]", 8, "0.0"),
    ("T_contact", "Kontaktzone\nmin. [°C]", 11.5, "0.0"),
]
DET_GROUPS = [("Geometrie", 2, 8), ("Ablauf und Kontakt", 9, 19), ("Kraft", 20, 21), ("Zeit", 22, 23),
              ("Änderung der Mitteltemperatur", 24, 27), ("Temperatur am Stichende", 28, 31)]


def _sheet_details(ws, D, title):
    rows, proc, blk = D["rows"], D["proc"], D["block"]
    ncol = len(DET_COLS)
    ws.sheet_view.showGridLines = False
    for j, (_k, _h, w, _f) in enumerate(DET_COLS, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    put(ws, "A1", f"{title}: alle Werte je Stich", font(14, True))
    ws.row_dimensions[1].height = 24
    put(ws, "A2", "Die Änderung der Mitteltemperatur je Stich setzt sich aus Umformwärme, Wärmeabfuhr in die Sättel "
                  "und Abkühlung an Luft zusammen. Weitere Erläuterungen auf dem Blatt „Parameter“.",
        font(9, color=MUTED))
    r_grp, r_head = 4, 5
    for name, c1, c2 in DET_GROUPS:
        merge_put(ws, r_grp, c1, r_grp, c2, name, f=font(9, True, MUTED),
                  al=Alignment(horizontal="center", vertical="bottom"),
                  br=Border(bottom=side("thin", LINE_STRONG)))
    for j, (_k, head, _w, _f) in enumerate(DET_COLS, start=1):
        put(ws, (r_head, j), head, font(9, True), al=Alignment(horizontal="center", vertical="center", wrap_text=True),
            fl=fill(PANEL2), br=Border(bottom=side("medium", INK)))
    ws.row_dimensions[r_head].height = 30
    r = r_head
    st = D["start"]
    start_vals = dict(nr=0, form=SHAPES[blk["shape"]][0], d=blk["d"], L=blk["L"], A_cm2=blk["A"] / 100,
                      T_core=st[0], T_surf=st[1], T_mean=st[2], t_end_min=D.get("t_transfer", 0.0) / 60.0)
    all_rows = [(start_vals, {"form", "d", "L"}, {}, True)]
    for row in rows:
        v = dict(row)
        v.update(form=SHAPES[row["shape"]][0], A_cm2=row["A"] / 100, A_K_cm2=row["A_K"] / 100,
                 sB_h=row["sB"] / row["h_prev"], t_end_min=row["t_end"] / 60.0, dT_die_s=-row["dT_die"],
                 dT_air=row["dT_mean"] - row["dT_def"] + row["dT_die"])
        inputs = {"form", "d" if row["mode"] == "dim" else "L"}
        if row.get("nrot_manual"):
            inputs.add("n_rot")
        if row.get("bk_manual"):
            inputs.add("bc")
        all_rows.append((v, inputs, _flags(row, proc), False))
    for i, (vals, inputs, flags, is_start) in enumerate(all_rows):
        r += 1
        for j, (key, _h, _w, fmt) in enumerate(DET_COLS, start=1):
            v = vals.get(key)
            col = INPUT if key in inputs else (MUTED if is_start else INK)
            bg = ZEBRA if (not is_start and i % 2 == 0) else None
            bold = key == "nr"
            if key in flags:
                col, bg, bold = (ERR, ERR_BG, True) if flags[key] == "err" else (WARN, WARN_BG, True)
            al = Alignment(horizontal="center" if key == "nr" else ("left" if key == "form" else "right"),
                           vertical="center", indent=1 if key == "form" else 0)
            put(ws, (r, j), v, font(10, bold, col), fmt=fmt if v is not None else None, al=al,
                fl=fill(bg) if bg else None, br=Border(bottom=side("thin", LINE)))
        ws.row_dimensions[r].height = 19
    for j in range(1, ncol + 1):
        ws.cell(row=r, column=j).border = Border(bottom=side("medium", INK))
    r += 2
    put(ws, (r, 1), "Blau: Eingabe im Programm, schwarz: berechnet. Stich 0 ist der Rohblock nach dem Transport zur "
                    "Presse. Kontaktfläche je Sattel = Bisslänge × Kontaktbreite.", font(8.5, color=MUTED))
    ws.freeze_panes = ws.cell(row=r_head + 1, column=3)
    setup_print(ws, True, ncol, r, title)
    ws.print_title_rows = f"{r_grp}:{r_head}"


# ---------------------------------------------------------------- Blatt 3: Parameter
def _sheet_parameter(ws, D, title):
    mat, blk, proc, kf = D["mat"], D["block"], D["proc"], D["kf"]
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCD", (36, 14, 13, 62)):
        ws.column_dimensions[col].width = w
    put(ws, "A1", f"{title}: Parameter und Rechenmodell", font(14, True))
    ws.row_dimensions[1].height = 24
    put(ws, "A2", "Eingaben sind blau, feste Konstanten des Rechenmodells schwarz.", font(9, color=MUTED))
    r = 3
    a_mm2s = mat["k"] / (mat["rho"] * mat["cp"]) * 1e6
    E, K = "Eingabe", "Konstante im Modell"
    sections = [
        ("Werkstoff", [
            ("Bezeichnung", mat["name"], "", None, E),
            ("Dichte ρ", mat["rho"], "kg/m³", "#,##0", E),
            ("Wärmekapazität c_p", mat["cp"], "J/(kg K)", "#,##0", E),
            ("Wärmeleitfähigkeit λ", mat["k"], "W/(m K)", "0.0", E),
            ("Temperaturleitfähigkeit a = λ/(ρ c_p)", a_mm2s, "mm²/s", "0.0", "berechnet"),
        ]),
        ("Rohblock", [
            ("Form", SHAPES[blk["shape"]][0], "", None, E),
            ({"rund": "Durchmesser D", "vier": "Kantenlänge a", "acht": "Schlüsselweite SW"}[blk["shape"]],
             blk["d"], "mm", "0.0", E),
        ] + ([("Kantenradius r", blk["r"], "mm", "0.0", E)] if blk["shape"] != "rund" else []) + [
            ("Länge L", blk["L"], "mm", "#,##0", E),
            ("Querschnitt A", blk["A"] / 100, "cm²", "0.0", "berechnet"),
            ("Volumen V", blk["A"] * blk["L"] / 1e6, "dm³", "0.00", "berechnet"),
            ("Masse m", blk["mass"], "kg", "0.0", "berechnet"),
        ]),
        ("Ablauf", [
            ("Temperatur bei Ofenentnahme", proc["T_start"], "°C", "0", E),
            ("Transport zur Presse", proc["t_transfer"], "s", "0", E),
            ("Nebenzeit je Hub", proc["t_aux"], "s", "0.0", "Lüften, Vorschub, Drehen, Anfahren; auf dem Untersattel"),
            ("Pause zwischen Stichen", proc["t_pause"], "s", "0", E),
        ]),
        ("Presse und Sättel", [
            ("Hubgeschwindigkeit v", proc["v"], "mm/s", "0.0", E),
            ("Verweilzeit unter Last", proc["t_dwell"], "s", "0.00", E),
            ("Sattelbreite", proc["b_die"], "mm", "0", E),
            ("Bisslänge", {"min": "Sattelbreite", "const": "konstant", "sbh": "aus s_B/h"}[proc.get("bite_mode", "min")],
             "", None, "Vorschub = Bisslänge × A_vorher/A_nachher; Bisslänge höchstens Sattelbreite"),
            ("Bisslänge konstant", proc.get("sB_const"), "mm", "0", "nur bei konstanter Bisslänge"),
            ("Bisslänge s_B / Höhe h", proc.get("sB_rel"), "–", "0.00", "nur bei Bisslänge aus s_B/h"),
        ]),
        ("Wärmeübergang", [
            ("Satteltemperatur", proc["T_die"], "°C", "0", E),
            ("Kontaktleitwert h_K unter Last", proc["h_c"], "W/(m² K)", "#,##0", E),
            ("Kontaktleitwert beim Aufliegen", proc.get("h_rest", 0.0), "W/(m² K)", "#,##0",
             "Nebenzeit auf dem Untersattel (QForm: CoolingInTool)"),
            ("Wärmeleitfähigkeit Sattel λ", proc["lam_die"], "W/(m K)", "0.0", E),
            ("ρ·c Sattel", 3.6, "MJ/(m³ K)", "0.0", K + " (Warmarbeitsstahl)"),
            ("Wärmeübergang Luft h", proc["h_air"], "W/(m² K)", "0.0", E),
            ("Emissionsgrad ε", proc["eps"], "–", "0.00", E),
            ("Umgebungstemperatur", 20, "°C", "0", K),
        ]),
        ("Umformung und Fließspannung", _kf_rows(kf, E) + [
            ("Formänderungswirkungsgrad η_F", proc["eta_F"], "–", "0.00", E),
            ("Reibungszahl μ", proc["mu"], "–", "0.00", E),
            ("Taylor-Quinney-Koeffizient β", 0.9, "–", "0.00", K),
        ]),
        ("Grenzen", [
            ("Presskraft max.", proc["F_max"], "MN", "0.0", E),
            ("Mindesttemperatur Oberfläche", proc["T_min"], "°C", "0", E),
            ("Höchsttemperatur Kern", proc["T_max"], "°C", "0", E),
        ]),
    ]
    for name, items in sections:
        r += 1
        ws.row_dimensions[r].height = 8
        r += 1
        for j, h in enumerate((name, "Wert", "Einheit", "Hinweis"), start=1):
            put(ws, (r, j), h, font(10 if j == 1 else 9, True, INK if j == 1 else MUTED), fl=fill(PANEL2),
                al=Alignment(horizontal="right" if j == 2 else "left", vertical="center", indent=0 if j == 2 else 1),
                br=Border(bottom=side("medium", INK)))
        ws.row_dimensions[r].height = 20
        for lab, val, unit, fmt, note in items:
            r += 1
            is_input = note == E or (note not in (K, "berechnet") and not note.startswith(K))
            put(ws, (r, 1), lab, font(10), al=Alignment(horizontal="left", indent=1, vertical="center"),
                br=Border(bottom=side("thin", LINE)))
            put(ws, (r, 2), val, font(10, False, INPUT if is_input else INK), fmt=fmt,
                al=Alignment(horizontal="right", vertical="center"), br=Border(bottom=side("thin", LINE)))
            put(ws, (r, 3), unit, font(10, color=MUTED), al=Alignment(horizontal="left", indent=1, vertical="center"),
                br=Border(bottom=side("thin", LINE)))
            put(ws, (r, 4), "" if note == E else note, font(9, color=MUTED),
                al=Alignment(horizontal="left", indent=1, vertical="center"),
                br=Border(bottom=side("thin", LINE)))
            ws.row_dimensions[r].height = 17
    # Rechenmodell (neue Druckseite)
    r += 2
    ws.row_breaks.append(Break(id=r - 1))
    put(ws, (r, 1), "Rechenmodell", font(12, True))
    for j in range(1, 5):
        ws.cell(row=r, column=j).border = Border(bottom=side("medium", INK))
    ws.row_dimensions[r].height = 20
    for head, text in D.get("model", []):
        r += 1
        ws.row_dimensions[r].height = 8
        r += 1
        put(ws, (r, 1), head, font(10, True))
        r += 1
        merge_put(ws, r, 1, r, 4, text, f=font(9.5), al=Alignment(wrap_text=True, vertical="top"))
        # Zeilenhöhe grob nach Textlänge (Breite A:D ≈ 125 Zeichen bei 9,5 pt)
        lines = sum(max(1, math.ceil(len(p) / 118)) for p in text.split("\n"))
        ws.row_dimensions[r].height = 13.5 * lines + 4
    setup_print(ws, False, 4, r, title)


def _kf_rows(kf, E):
    if kf["type"] == "st":
        return [
            ("Fließspannungsmodell", "Sellars-Tegart", "", None, "σ = 1/α · arsinh[(Z/A)^(1/n)], Z = φ̇ · exp(Q/RT)"),
            ("α", kf["alpha"], "1/MPa", "0.0000", E),
            ("n", kf["n"], "–", "0.00", E),
            ("Q", kf["Q"], "kJ/mol", "0.0", E),
            ("ln A", kf["lnA"], "ln(1/s)", "0.000", E),
        ]
    out = [("Fließspannungsmodell", "Hensel-Spittel", "", None, "T in °C, φ̇ in 1/s, Ergebnis in MPa")]
    for key in ("A", "m1", "m2", "m3", "m4", "m5", "m7", "m8", "m9"):
        out.append((key, kf[key], "MPa" if key == "A" else "–", "0.000" if key == "A" else "0.000000", E))
    return out


# ---------------------------------------------------------------- Blatt 4: Diagrammdaten
def _sheet_chartdata(ws, D):
    s = D["series"]
    proc, rows = D["proc"], D["rows"]
    t, core, surf = s["t"], s["core"], s["surf"]
    idx = _decimate(t, surf)
    for col, w in zip("ABCDEFGHIJKLMNO", (10, 10, 11, 13, 3, 12, 14, 3, 11, 11, 11, 3, 8, 12, 12)):
        ws.column_dimensions[col].width = w
    put(ws, "A1", "Diagrammdaten", font(14, True))
    put(ws, "A2", f"Temperaturverlauf auf {len(idx)} Punkte ausgedünnt (Extremwerte bleiben erhalten); "
                  "die vollständige Zeitreihe enthält der CSV-Export.", font(9, color=MUTED))
    hdr_style = dict(f=font(9, True), fl=fill(PANEL2), br=Border(bottom=side("medium", INK)),
                     al=Alignment(horizontal="center", vertical="center", wrap_text=True))
    heads = {1: "Zeit\n[s]", 2: "Zeit\n[min]", 3: "Kern", 4: "Oberfläche (Mittel)", 6: "Zeit\n[min]",
             7: "Kontaktzone am Hubende", 9: "Zeit\n[min]", 10: "Mindesttemperatur", 11: "Höchsttemperatur Kern", 13: "Stich",
             14: "Presskraft\n[MN]", 15: "Grenze Presskraft"}
    r_head = 3
    for c, h in heads.items():
        put(ws, (r_head, c), h, **hdr_style)
    ws.row_dimensions[r_head].height = 30
    r0 = r_head + 1
    for k, i in enumerate(idx):
        r = r0 + k
        ws.cell(row=r, column=1, value=round(t[i], 2)).number_format = "0.0"
        ws.cell(row=r, column=2, value=round(t[i] / 60.0, 4)).number_format = "0.000"
        ws.cell(row=r, column=3, value=round(core[i], 2)).number_format = "0.0"
        ws.cell(row=r, column=4, value=round(surf[i], 2)).number_format = "0.0"
    t_last = r0 + len(idx) - 1
    cts = s["contacts"]
    for k, (tc, Tc) in enumerate(cts):
        ws.cell(row=r0 + k, column=6, value=round(tc / 60.0, 4)).number_format = "0.000"
        ws.cell(row=r0 + k, column=7, value=round(Tc, 2)).number_format = "0.0"
    c_last = r0 + len(cts) - 1
    t_max = t[-1] / 60.0
    for k, x in enumerate((0.0, t_max)):
        ws.cell(row=r0 + k, column=9, value=round(x, 4)).number_format = "0.000"
        ws.cell(row=r0 + k, column=10, value=proc["T_min"]).number_format = "0"
        ws.cell(row=r0 + k, column=11, value=proc["T_max"]).number_format = "0"
    for k, row in enumerate(rows):
        ws.cell(row=r0 + k, column=13, value=k + 1)
        ws.cell(row=r0 + k, column=14, value=round(row["F"], 4)).number_format = "0.000"
        ws.cell(row=r0 + k, column=15, value=proc["F_max"]).number_format = "0.0"
    for row in ws.iter_rows(min_row=r0, max_row=max(t_last, c_last, r0 + len(rows) - 1), max_col=15):
        for c in row:
            if c.value is not None:
                c.font = font(9)
    ws.freeze_panes = "A4"
    # Achsenbereich: Kurven, Kontaktpunkte und Grenzen, sofern sie in der Nähe liegen
    vals = list(core) + list(surf) + [c[1] for c in cts]
    lo, hi = min(vals), max(vals)
    if proc["T_min"] > lo - 60:
        lo = min(lo, proc["T_min"])
    if proc["T_max"] < hi + 40:
        hi = max(hi, proc["T_max"])
    pad = max(hi - lo, 20.0) * 0.05
    # Grenzlinien außerhalb des Bereichs nicht zeichnen: Zellen leeren
    for col, T in ((10, proc["T_min"]), (11, proc["T_max"])):
        if not (lo - pad <= T <= hi + pad):
            ws.cell(row=r0, column=col).value = None
            ws.cell(row=r0 + 1, column=col).value = None
    setup_print(ws, False, 15, max(t_last, c_last), "Diagrammdaten")
    return dict(sheet=ws, t_first=r0, t_last=t_last, c_first=r0, c_last=c_last, g_first=r0, f_first=r0,
                f_last=r0 + len(rows) - 1, T_lo=lo - pad, T_hi=hi + pad, t_max=t_max)
